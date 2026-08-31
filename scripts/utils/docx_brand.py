# -*- coding: utf-8 -*-
"""Shared Taleemabad-branded DOCX builder.

Extracted from scripts/case_studies/make_rm_case_study_doc.py so any internal or
candidate-facing Word document can use the same locked layout:
logo header, Quicksand, navy/blue palette, branded tables, callout boxes.

Spec: memory/case_studies_smg_gm_from_hog_2026_07_31.md (layout section).
"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = r"c:\Agent Coco\assets\logo_taleemabad.png"
FONT = "Quicksand"

BLUE = RGBColor(0x3C, 0x78, 0xD8)
NAVY = RGBColor(0x2F, 0x4F, 0xA2)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FILL_HEAD = "2F4FA2"
FILL_ALT = "F2F6FC"
FILL_CALL = "EEF3FB"
RULE_CLR = "C9D6EE"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def style(run, size, color, bold, italic, underline):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    run.underline = underline
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(attr), FONT)


def add_runs(p, text, size=11, color=BLACK, italic=False, underline=False, bold_all=False):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            style(r, size, color, bold_all, italic, underline)
        r = p.add_run(m.group(1))
        style(r, size, color, True, italic, underline)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        style(r, size, color, bold_all, italic, underline)


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexfill)
    tcPr.append(shd)


def cell_borders(cell, color=RULE_CLR, sz=6, edges=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in edges:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def clear_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tcPr.append(borders)


def accent_bar(cell, color="3C78D8", sz=24):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ("top", "bottom", "right"):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(sz))
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), color)
    borders.append(left)
    tcPr.append(borders)


def cell_margins(cell, top=80, bottom=80, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement('w:' + name)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tcPr.append(mar)


def hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), RULE_CLR)
    borders.append(bottom)
    pPr.append(borders)


def make_table(doc, rows, widths, size=10.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Inches(widths[ci])
            cell_margins(cell)
            cell_borders(cell)
            if ri == 0:
                shade(cell, FILL_HEAD)
            elif ri % 2 == 0:
                shade(cell, FILL_ALT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, val, size=size,
                     color=WHITE if ri == 0 else BLACK,
                     bold_all=(ri == 0))
    return t


def make_callout(doc, lines, color="3C78D8", fill=FILL_CALL, width=6.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Inches(width)
    cell_margins(cell, top=120, bottom=120, left=180, right=180)
    accent_bar(cell, color)
    shade(cell, fill)
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for i, (kind, text) in enumerate(lines):
        p = cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0 if i == 0 else 5)
        pf.space_after = Pt(0)
        if kind == "who":
            add_runs(p, text, size=10.5, color=NAVY, bold_all=True)
        elif kind == "plain":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, text, size=10.5)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, text, size=11, italic=True)
    return t


def banner(doc, left_text, right_text):
    """Navy bar: label left, tag right."""
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    left, right = t.cell(0, 0), t.cell(0, 1)
    left.width = Inches(4.9)
    right.width = Inches(1.6)
    for cell in (left, right):
        shade(cell, FILL_HEAD)
        clear_borders(cell)
        cell_margins(cell, top=90, bottom=90, left=160, right=160)
    lp = left.paragraphs[0]
    lp.paragraph_format.space_before = Pt(0); lp.paragraph_format.space_after = Pt(0)
    add_runs(lp, left_text, size=12.5, color=WHITE, bold_all=True)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = Pt(0); rp.paragraph_format.space_after = Pt(0)
    add_runs(rp, right_text, size=10.5, color=WHITE, bold_all=True)
    return t


def build(spec, path, logo=LOGO, footer_note=None):
    """Render a spec (list of (kind, ...) tuples) to a branded DOCX."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.add_run().add_picture(logo, width=Inches(1.15))
        if footer_note:
            fp = section.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(fp, footer_note, size=8.5, color=GREY, italic=True)

    for item in spec:
        kind = item[0]

        if kind == "rule":
            hrule(doc); continue
        if kind == "pagebreak":
            doc.add_page_break(); continue
        if kind == "table":
            make_table(doc, item[1], item[2], *(item[3:])); continue
        if kind == "callout":
            make_callout(doc, item[1], *(item[2:])); continue
        if kind == "banner":
            banner(doc, item[1], item[2] if len(item) > 2 else ""); continue
        if kind == "spacer":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(item[1] if len(item) > 1 else 6)
            continue

        text = item[1] if len(item) > 1 else ""
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(6)
        if kind == "title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(18); pf.space_after = Pt(4)
            add_runs(p, text, size=17, bold_all=True)
        elif kind == "subtitle":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after = Pt(16)
            add_runs(p, text, size=15, color=BLUE, bold_all=True)
        elif kind == "meta":
            pf.space_after = Pt(2)
            add_runs(p, text, size=10.5)
        elif kind == "part":
            pf.space_before = Pt(18); pf.space_after = Pt(2)
            add_runs(p, text, size=15, color=NAVY, bold_all=True)
        elif kind == "h1":
            pf.space_before = Pt(16); pf.space_after = Pt(8)
            add_runs(p, text, size=15, bold_all=True)
        elif kind == "h2":
            pf.space_before = Pt(14); pf.space_after = Pt(6)
            add_runs(p, text, size=12.5, color=BLUE, bold_all=True)
        elif kind == "h3":
            pf.space_before = Pt(8); pf.space_after = Pt(4)
            add_runs(p, text, bold_all=True)
        elif kind == "h3u":
            pf.space_before = Pt(8); pf.space_after = Pt(4)
            add_runs(p, text, bold_all=True, underline=True)
        elif kind == "para":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, text)
        elif kind == "caption":
            pf.space_before = Pt(4)
            add_runs(p, text, size=9.5, color=GREY, italic=True)
        elif kind == "bullet":
            p.style = doc.styles["List Bullet"]
            pf.space_after = Pt(3)
            add_runs(p, text)
        elif kind == "sbullet":                     # compact bullet
            p.style = doc.styles["List Bullet"]
            pf.space_after = Pt(1.5)
            add_runs(p, text, size=9.5)
        elif kind == "spara":                       # compact body
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_after = Pt(4)
            add_runs(p, text, size=9.5)
        elif kind == "sh2":                         # compact blue heading
            pf.space_before = Pt(7); pf.space_after = Pt(2)
            add_runs(p, text, size=10.5, color=BLUE, bold_all=True)
        elif kind == "bullet2":
            p.style = doc.styles["List Bullet 2"]
            pf.space_after = Pt(3)
            add_runs(p, text)
        elif kind == "num":
            p.style = doc.styles["List Number"]
            pf.space_after = Pt(3)
            add_runs(p, text)
        elif kind == "footer":
            pf.space_before = Pt(14); pf.space_after = Pt(0)
            add_runs(p, text, italic=True)
    doc.save(path)
    print("built", path)
    return path


def upload_as_gdoc(path, title, file_id=None,
                   token=r"c:\Agent Coco\.claude\config\token_sheets_broad.json"):
    """Create a native Google Doc from a DOCX, or overwrite an existing one.

    Pass file_id to update in place so the shared link stays valid.
    """
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build as gbuild
    from googleapiclient.http import MediaFileUpload
    creds = Credentials.from_authorized_user_file(token)
    drive = gbuild("drive", "v3", credentials=creds)
    media = MediaFileUpload(
        path,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    if file_id:
        f = drive.files().update(fileId=file_id, body={"name": title},
                                 media_body=media,
                                 fields="id, webViewLink").execute()
        print("updated in place:", title)
    else:
        f = drive.files().create(
            body={"name": title, "mimeType": "application/vnd.google-apps.document"},
            media_body=media, fields="id, webViewLink",
        ).execute()
        print(title)
    print("  ->", f["webViewLink"])
    return f["webViewLink"]
