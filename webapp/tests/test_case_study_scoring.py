"""The locked case-study scoring rules, tested without a model or a database.

The zero is the point: anchoring at 1 floors every dimension at 20% of its weight,
which is the defect that put all 25 RM case studies above the bar (CLAUDE.md Rule 27).
"""

from __future__ import annotations

import datetime as dt

import pytest
from fastapi import HTTPException

from webapp.services.case_study_scoring import (
    DIMENSIONS,
    SCORES,
    CaseStudyScoringError,
    band,
    score_submission,
    validate_scores,
    weighted_total,
)


def _all(n):
    return {d["key"]: n for d in DIMENSIONS}


def test_weights_are_the_locked_six_and_sum_to_100():
    assert [d["weight"] for d in DIMENSIONS] == [20, 25, 20, 15, 10, 10]
    assert sum(d["weight"] for d in DIMENSIONS) == 100
    assert len(DIMENSIONS) == 6


def test_the_scale_has_a_real_zero():
    assert SCORES == (0, 1, 2, 3, 4, 5)
    # THE defect this rubric change exists to remove: all-minimum must be 0, not 20.
    assert weighted_total(_all(0)) == 0.0


def test_conversion_is_linear_on_the_full_scale():
    assert weighted_total(_all(5)) == 100.0
    assert weighted_total(_all(4)) == 80.0
    assert weighted_total(_all(3)) == 60.0
    assert weighted_total(_all(2)) == 40.0
    assert weighted_total(_all(1)) == 20.0


def test_weighting_is_per_dimension_not_flat():
    # Only Execution specificity (weight 25) at full marks.
    scores = _all(0)
    scores["execution_specificity"] = 5
    assert weighted_total(scores) == 25.0


def test_bands_match_the_locked_thresholds():
    assert band(80.0, []) == "strong_yes"
    assert band(95.0, []) == "strong_yes"
    assert band(79.9, []) == "yes"
    assert band(65.0, []) == "yes"
    assert band(64.9, []) == "borderline"
    assert band(50.0, []) == "borderline"
    assert band(49.9, []) == "no"
    assert band(0.0, []) == "no"


def test_fabricated_data_disqualifies_regardless_of_total():
    assert band(100.0, ["fabricated_data"]) == "disqualified"
    assert band(92.0, ["fabricated_data", "undisclosed_ai"]) == "disqualified"
    # A serious-but-not-disqualifying flag does NOT override the band.
    assert band(92.0, ["undisclosed_ai"]) == "strong_yes"


def test_benchmark_table_is_in_the_coco_schema():
    from webapp.models import EvalBenchmark

    assert EvalBenchmark.__table__.schema == "coco"


def test_benchmark_carries_a_qa_gate():
    from webapp.models import EvalBenchmark

    cols = {c.name for c in EvalBenchmark.__table__.columns}
    # Rule 0 is enforced by requiring these before a run may start.
    assert {"qa_approved_at", "qa_approved_by", "status"} <= cols


# ── Task 6: evaluation storage (webapp/models.py::CaseStudyEvaluation) ─────
#
# A scored evaluation persists what score_submission returned so
# GET /api/case-studies/evaluations/{id} can read it back later -- see
# webapp/routers/case_studies.py.


def test_case_study_evaluation_table_is_in_the_coco_schema():
    from webapp.models import CaseStudyEvaluation

    assert CaseStudyEvaluation.__table__.schema == "coco"


def test_case_study_evaluation_carries_the_benchmark_it_was_scored_against():
    from webapp.models import CaseStudyEvaluation

    cols = {c.name for c in CaseStudyEvaluation.__table__.columns}
    # "which answer key produced this total" must always be answerable from
    # the row alone -- never just "the job's benchmark" (a job can carry more
    # than one benchmark row over time).
    assert {"benchmark_id", "job_id", "application_id"} <= cols
    assert {"scores", "evidence", "flags", "total", "band"} <= cols


def test_case_study_evaluation_references_eval_benchmarks_by_foreign_key():
    from webapp.models import CaseStudyEvaluation

    fks = {
        fk.target_fullname
        for col in CaseStudyEvaluation.__table__.columns
        for fk in col.foreign_keys
    }
    assert "coco.eval_benchmarks.id" in fks


def test_validate_rejects_a_bad_score_set():
    with pytest.raises(CaseStudyScoringError):
        validate_scores({d["key"]: 3 for d in DIMENSIONS[:5]})     # missing one
    with pytest.raises(CaseStudyScoringError):
        bad = _all(3); bad["data_judgment"] = 6
        validate_scores(bad)                                        # out of range
    with pytest.raises(CaseStudyScoringError):
        bad = _all(3); bad["not_a_dimension"] = 3
        validate_scores(bad)                                        # unknown key
    with pytest.raises(CaseStudyScoringError):
        bad = _all(3); bad["data_judgment"] = 2.5
        validate_scores(bad)                                        # not an integer
    validate_scores(_all(0))                                        # zero is VALID


# ── Task 5: the model call (webapp/prompts/case_study_prompt.py) ───────────
#
# These monkeypatch _call_model, so no network call is ever made here.


def _good_evidence():
    return {d["key"]: f"See the {d['label']} section of the submission." for d in DIMENSIONS}


def _good_response(scores=None, flags=None):
    return {
        "scores": scores if scores is not None else _all(3),
        "evidence": _good_evidence(),
        "flags": flags or [],
    }


def test_total_and_band_are_computed_not_taken_from_the_model(monkeypatch):
    """A model returning total: 95 on a score set worth 40 must be overruled."""
    from webapp.services import case_study_scoring as css

    scores = _all(2)  # weighted_total(_all(2)) == 40.0, per the pure-rules test above
    payload = {
        "scores": scores,
        "evidence": _good_evidence(),
        "flags": [],
        "total": 95,            # the model lying
        "band": "strong_yes",   # the model lying
    }
    monkeypatch.setattr(css, "_call_model", lambda **kw: (payload, "test-model"))

    out = score_submission(
        corpus="submission text", benchmark_body="benchmark text",
        candidate_name="Zara Khan", role="Growth Manager",
    )
    assert out["total"] == 40.0        # computed, never the model's 95
    assert out["band"] == "no"         # 40 is below the 50 borderline floor


def test_fabricated_data_flag_disqualifies_even_at_a_100_total(monkeypatch):
    from webapp.services import case_study_scoring as css

    payload = _good_response(scores=_all(5), flags=["fabricated_data"])
    monkeypatch.setattr(css, "_call_model", lambda **kw: (payload, "test-model"))

    out = score_submission(
        corpus="submission text", benchmark_body="benchmark text",
        candidate_name="Zara Khan", role="Growth Manager",
    )
    assert out["total"] == 100.0
    assert out["band"] == "disqualified"


def test_blank_evidence_citation_is_rejected(monkeypatch):
    """Rule 1: a dimension score with no quoted line, slide number or figure
    behind it is not a score. Both attempts return the same blank citation,
    so this must raise rather than be silently repaired."""
    from webapp.services import case_study_scoring as css

    bad_evidence = _good_evidence()
    bad_evidence["data_judgment"] = "   "  # whitespace only
    payload = {"scores": _all(3), "evidence": bad_evidence, "flags": []}
    monkeypatch.setattr(css, "_call_model", lambda **kw: (payload, "test-model"))

    with pytest.raises(CaseStudyScoringError):
        score_submission(
            corpus="submission text", benchmark_body="benchmark text",
            candidate_name="Zara Khan", role="Growth Manager",
        )


def test_a_bare_value_error_from_the_model_call_is_retried_and_can_succeed(monkeypatch):
    """_call_model can raise a plain ValueError (drafting._parse_json's "LLM
    did not return parseable JSON") before any CaseStudyScoringError territory
    is reached. That must be retried exactly like a wrong-shape response."""
    from webapp.services import case_study_scoring as css

    good = _good_response()
    calls = []

    def flaky(**kw):
        calls.append(kw)
        if len(calls) == 1:
            raise ValueError("Expecting value: line 1 column 1")
        return good, "test-model"

    monkeypatch.setattr(css, "_call_model", flaky)
    out = score_submission(
        corpus="submission text", benchmark_body="benchmark text",
        candidate_name="Zara Khan", role="Growth Manager",
    )
    assert out["model"] == "test-model"
    assert out["total"] == 60.0
    assert len(calls) == 2


def test_two_bare_value_errors_raise_case_study_scoring_error_not_bare_value_error(monkeypatch):
    """Two consecutive parse failures must still surface as a single,
    handleable exception type, and the stub must never be called a third
    time -- the retry budget is exactly one retry."""
    from webapp.services import case_study_scoring as css

    calls = []

    def always_broken(**kw):
        calls.append(kw)
        raise ValueError("Expecting value: line 1 column 1")

    monkeypatch.setattr(css, "_call_model", always_broken)
    with pytest.raises(CaseStudyScoringError) as excinfo:
        score_submission(
            corpus="submission text", benchmark_body="benchmark text",
            candidate_name="Zara Khan", role="Growth Manager",
        )
    assert type(excinfo.value) is CaseStudyScoringError
    assert len(calls) == 2


def test_score_submission_returns_the_full_locked_shape(monkeypatch):
    from webapp.services import case_study_scoring as css

    payload = _good_response(scores=_all(4))
    monkeypatch.setattr(css, "_call_model", lambda **kw: (payload, "test-model"))

    out = score_submission(
        corpus="submission text", benchmark_body="benchmark text",
        candidate_name="Zara Khan", role="Growth Manager",
    )
    assert set(out) == {"scores", "evidence", "flags", "total", "band", "model"}
    assert out["scores"] == _all(4)
    assert out["evidence"] == _good_evidence()
    assert out["flags"] == []
    assert out["total"] == 80.0
    assert out["band"] == "strong_yes"
    assert out["model"] == "test-model"


def test_unknown_flag_from_the_model_is_rejected(monkeypatch):
    """The flag vocabulary is closed. A model inventing its own flag name
    must not silently ride through into the band computation."""
    from webapp.services import case_study_scoring as css

    payload = _good_response(flags=["not_a_real_flag"])
    monkeypatch.setattr(css, "_call_model", lambda **kw: (payload, "test-model"))

    with pytest.raises(CaseStudyScoringError):
        score_submission(
            corpus="submission text", benchmark_body="benchmark text",
            candidate_name="Zara Khan", role="Growth Manager",
        )


def test_case_study_prompt_builds_without_a_model():
    """system_prompt()/build_user_prompt() are pure string assembly and read
    the locked rubric file at call time -- proving they succeed (and carry
    the three numbered rules) needs no model and no network."""
    from webapp.prompts import case_study_prompt

    system = case_study_prompt.system_prompt()
    assert "RULE 1" in system and "RULE 2" in system and "RULE 3" in system
    assert "fabricated_data" in system
    for d in DIMENSIONS:
        assert d["key"] in system

    user = case_study_prompt.build_user_prompt(
        corpus="the submission body",
        benchmark_body="the benchmark body",
        candidate_name="Zara Khan",
        role="Growth Manager",
    )
    assert "the submission body" in user
    assert "the benchmark body" in user
    assert "Zara Khan" in user


# ── Submission retrieval (webapp/services/submissions.py) ──────────────────
#
# The core rule: an unreadable submission is REFUSED, never scored as a weak
# one (mirrors cv_text.CVUnreadable). Every fetcher here is exercised through
# an injected fake -- no test in this module opens a socket, a mailbox, or
# Drive. A test that reaches the network is not a unit test.

from webapp.services.submissions import (  # noqa: E402
    MIN_USABLE_CHARS,
    RawAttachment,
    SubmissionContext,
    SubmissionUnreadable,
    _drive_file_id,
    _extract_links,
    corpus_for,
    markaz_case_study_api,
    markaz_linked_documents,
)


def _ctx(**overrides):
    base = dict(application_id=4242, candidate_name="Zara Khan",
                candidate_email="zara@example.com", job_title="Growth Manager")
    base.update(overrides)
    return SubmissionContext(**base)


def test_corpus_for_refuses_when_no_fetcher_finds_anything():
    """No attachments at all from any channel -> refuse, not an empty corpus."""
    empty_fetcher = lambda ctx: []  # noqa: E731
    with pytest.raises(SubmissionUnreadable):
        corpus_for(4242, context=_ctx(), fetchers=[empty_fetcher, empty_fetcher])


def test_corpus_for_refuses_below_the_usable_floor():
    """A source that DOES yield content, but under MIN_USABLE_CHARS, still
    refuses -- a thin scrap is treated the same as nothing, never scored."""
    thin = RawAttachment(origin="Gmail attachment: note.txt", filename="note.txt",
                          content=b"Not much here.")

    def fetcher(ctx):
        return [thin]

    with pytest.raises(SubmissionUnreadable) as exc_info:
        corpus_for(4242, context=_ctx(), fetchers=[fetcher])
    assert "4242" in str(exc_info.value)


def test_corpus_for_never_returns_usable_false():
    """The dict shape carries `usable: bool`, but the contract is refuse-or-
    succeed: corpus_for never hands back a dict with usable=False, mirroring
    cv_text.extract's raise-don't-degrade pattern."""
    thin = RawAttachment(origin="x", filename="x.txt", content=b"short")
    with pytest.raises(SubmissionUnreadable):
        corpus_for(4242, context=_ctx(), fetchers=[lambda ctx: [thin]])


def test_corpus_for_succeeds_and_records_every_source_actually_read():
    long_text = ("The candidate proposed a three-phase rollout for the growth "
                 "flywheel, sequencing partner recruitment before budget commitments. ") * 5
    assert len(long_text) >= 400
    att_a = RawAttachment(origin="Gmail attachment: case-study-4242-word-1.docx",
                           filename="submission.txt", content=long_text.encode("utf-8"))

    def mailbox_fetcher(ctx):
        return [att_a]

    def empty_fetcher(ctx):
        return []

    result = corpus_for(4242, context=_ctx(), fetchers=[mailbox_fetcher, empty_fetcher])

    assert result["usable"] is True
    assert result["chars"] == len(result["text"])
    assert result["chars"] >= 400
    assert result["sources"] == ["Gmail attachment: case-study-4242-word-1.docx"]
    assert "three-phase rollout" in result["text"]


def test_corpus_for_combines_multiple_sources_and_lists_each():
    part1 = "A" * 250
    part2 = "B" * 250
    att1 = RawAttachment(origin="Gmail attachment: part1.txt", filename="p1.txt",
                          content=part1.encode())
    att2 = RawAttachment(origin="Drive: https://drive.google.com/file/d/abc/view",
                          filename="p2.txt", content=part2.encode())

    result = corpus_for(
        4242, context=_ctx(),
        fetchers=[lambda ctx: [att1], lambda ctx: [att2]],
    )
    assert result["sources"] == [
        "Gmail attachment: part1.txt",
        "Drive: https://drive.google.com/file/d/abc/view",
    ]
    assert part1 in result["text"] and part2 in result["text"]


def test_corpus_for_skips_a_source_that_yields_nothing_but_keeps_trying():
    """A fetcher that raises (network error on one channel) must not abort the
    others -- it tries the next source, exactly like cv_text tries the next
    parser."""
    good = RawAttachment(origin="Gmail attachment: real.txt", filename="real.txt",
                          content=(b"C" * 500))

    def broken_fetcher(ctx):
        raise ConnectionError("mailbox unreachable")

    result = corpus_for(4242, context=_ctx(), fetchers=[broken_fetcher, lambda ctx: [good]])
    assert result["usable"] is True
    assert result["sources"] == ["Gmail attachment: real.txt"]


def test_corpus_for_needs_context_or_db():
    with pytest.raises(ValueError):
        corpus_for(4242)


def test_corpus_for_delegates_docx_extraction_to_fetch_submission_corpora():
    """Proves real delegation, not a reimplementation: a genuine .docx built
    with python-docx (same library cv_text.py uses) must come back as text via
    the shared fetch_submission_corpora.extract dispatch."""
    import io as _io

    import docx

    document = docx.Document()
    document.add_paragraph(
        "Our recommendation is to sequence the Lahore pilot before Karachi, "
        "because the partner data only supports one cohort at a time and the "
        "growth team cannot staff two onboarding tracks simultaneously this term."
    )
    document.add_paragraph(
        "Risks: the Karachi partner has a harder deadline, so a delay there "
        "carries more downside than a delay in Lahore, and we have built in a "
        "two-week buffer before the Karachi kickoff to absorb any slippage "
        "from the Lahore rollout finishing late."
    )
    buf = _io.BytesIO()
    document.save(buf)
    content = buf.getvalue()
    assert len(content) > 0

    att = RawAttachment(origin="Gmail attachment: plan.docx", filename="plan.docx", content=content)
    result = corpus_for(4242, context=_ctx(), fetchers=[lambda ctx: [att]])
    assert "Lahore pilot before Karachi" in result["text"]


def test_corpus_for_excludes_a_failed_extraction_sentinel_from_sources_and_chars(monkeypatch):
    """fetch_submission_corpora.extract() swallows its own exceptions and
    returns the truthy "[EXTRACT FAILED ...]" sentinel string instead of
    raising (documented in its own module). That string must never be treated
    as read text: not folded into the corpus, not counted toward
    extracted_chars (the readability floor), and not recorded in sources as a
    verified origin. Reproduces the exact production bug: with no
    python-pptx/openpyxl/PyMuPDF installed, a candidate's .pptx/.xlsx/.pdf
    extracted to this sentinel and was scored as if it were their real work,
    while the .docx they also submitted looked like the sole source."""
    from webapp.services import submissions as submissions_module

    good_text = "D" * 500
    good = RawAttachment(origin="Gmail attachment: writeup.docx", filename="writeup.docx",
                          content=b"irrelevant, extraction is faked below")
    bad = RawAttachment(origin="Gmail attachment: deck.pptx", filename="deck.pptx",
                         content=b"irrelevant, extraction is faked below")

    def fake_extract(path):
        if path.endswith(".pptx"):
            return "[EXTRACT FAILED ModuleNotFoundError: No module named 'pptx']"
        return good_text

    monkeypatch.setattr(submissions_module.extraction, "extract", fake_extract)

    result = corpus_for(4242, context=_ctx(), fetchers=[lambda ctx: [good, bad]])

    assert result["sources"] == ["Gmail attachment: writeup.docx"]
    assert "EXTRACT FAILED" not in result["text"]
    # "chars" legitimately includes the "===== origin =====" provenance header
    # around the one real source (see test_corpus_for_chars_field_still_
    # reflects_the_full_combined_text), so it is a little more than the raw
    # text -- what it must NOT include is anything from the failed .pptx.
    assert result["chars"] >= len(good_text)
    assert result["chars"] < len(good_text) + 100


def test_corpus_for_excludes_a_failed_extraction_sentinel_from_the_readability_floor(monkeypatch):
    """The sharper version of the test above: a THIN real source (below
    MIN_USABLE_CHARS on its own) alongside a LONG failed-extraction sentinel.
    If the sentinel were wrongly counted toward extracted_chars (the exact
    production bug -- a truthy error string satisfied
    `if extracted and extracted.strip():`), the combined length would clear
    MIN_USABLE_CHARS and the submission would be scored on a source that
    never actually extracted. With the fix, only the thin genuine text counts
    toward the floor, so this case correctly still refuses."""
    from webapp.services import submissions as submissions_module

    thin_good_text = "D" * 200  # real text, but alone under MIN_USABLE_CHARS (400)
    long_sentinel = "[EXTRACT FAILED ModuleNotFoundError: " + ("x" * 300) + "]"
    assert len(thin_good_text) < submissions_module.MIN_USABLE_CHARS
    assert len(thin_good_text) + len(long_sentinel) >= submissions_module.MIN_USABLE_CHARS

    good = RawAttachment(origin="Gmail attachment: writeup.docx", filename="writeup.docx",
                          content=b"irrelevant, extraction is faked below")
    bad = RawAttachment(origin="Gmail attachment: deck.pptx", filename="deck.pptx",
                         content=b"irrelevant, extraction is faked below")

    def fake_extract(path):
        return long_sentinel if path.endswith(".pptx") else thin_good_text

    monkeypatch.setattr(submissions_module.extraction, "extract", fake_extract)

    with pytest.raises(submissions_module.SubmissionUnreadable):
        corpus_for(4242, context=_ctx(), fetchers=[lambda ctx: [good, bad]])


def test_corpus_for_refuses_when_every_source_fails_extraction(monkeypatch):
    """A submission whose ONLY sources all fail extraction must be refused
    (SubmissionUnreadable), never scored on the sentinel text -- and the
    refusal must name the failure (it "lands in problems"), not report the
    generic "no source returned any content", which would hide that a file
    WAS found and read, just not successfully."""
    from webapp.services import submissions as submissions_module

    bad = RawAttachment(origin="Gmail attachment: deck.pptx", filename="deck.pptx",
                         content=b"irrelevant, extraction is faked below")

    def fake_extract(path):
        return "[EXTRACT FAILED ModuleNotFoundError: No module named 'pptx']"

    monkeypatch.setattr(submissions_module.extraction, "extract", fake_extract)

    with pytest.raises(submissions_module.SubmissionUnreadable) as exc_info:
        corpus_for(4242, context=_ctx(), fetchers=[lambda ctx: [bad]])

    detail = str(exc_info.value)
    assert "EXTRACT FAILED" in detail
    assert "ModuleNotFoundError" in detail
    assert "deck.pptx" in detail


def test_extract_links_finds_drive_and_markaz_upload_urls_only():
    blob = (
        "Please see my case study here: https://drive.google.com/file/d/1AbC-XyZ/view?usp=sharing "
        "and the workbook at https://markaz.taleemabad.com/uploads/case-study-99.xlsx. "
        "Also see my portfolio at https://myportfolio.example.com/work (not a submission link)."
    )
    links = _extract_links(blob)
    assert "https://drive.google.com/file/d/1AbC-XyZ/view?usp=sharing" in links
    assert "https://markaz.taleemabad.com/uploads/case-study-99.xlsx" in links
    assert not any("myportfolio" in link for link in links)


def test_extract_links_on_empty_or_none_is_empty():
    assert _extract_links("") == []
    assert _extract_links(None) == []


def test_drive_file_id_parses_the_common_share_url_shapes():
    real_length_id = "1AbC-XyZ9876543210abcdef"  # realistic Drive file ID length
    assert _drive_file_id(f"https://drive.google.com/file/d/{real_length_id}/view") == real_length_id
    assert _drive_file_id(f"https://drive.google.com/open?id={real_length_id}") == real_length_id
    assert _drive_file_id("https://example.com/not-a-drive-link") is None


def test_markaz_linked_documents_skips_the_spa_shell_trap():
    """markaz.taleemabad.com/uploads/... answers HTTP 200 with the SPA's own
    index.html for a missing file -- content-type must be checked, not status."""

    class FakeResponse:
        status_code = 200
        headers = {"content-type": "text/html; charset=utf-8"}
        content = b"<!doctype html><html>...</html>"

    def fake_get(url, timeout=15):
        return FakeResponse()

    ctx = _ctx(links=("https://markaz.taleemabad.com/uploads/case-study-99.docx",))
    result = markaz_linked_documents(ctx, http_get=fake_get)
    assert result == []


def test_markaz_linked_documents_accepts_a_real_content_type():
    class FakeResponse:
        status_code = 200
        headers = {
            "content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        }
        content = b"real bytes"

    def fake_get(url, timeout=15):
        return FakeResponse()

    ctx = _ctx(links=("https://markaz.taleemabad.com/uploads/case-study-99.docx",))
    result = markaz_linked_documents(ctx, http_get=fake_get)
    assert len(result) == 1
    assert result[0].content == b"real bytes"


def test_markaz_linked_documents_uses_drive_download_for_drive_urls():
    calls = []
    real_length_id = "1AbC-XyZ9876543210abcdef"

    def fake_drive_download(file_id):
        calls.append(file_id)
        return "Case Study - Zara Khan.docx", b"drive bytes"

    ctx = _ctx(links=(f"https://drive.google.com/file/d/{real_length_id}/view",))
    result = markaz_linked_documents(ctx, drive_download=fake_drive_download)
    assert calls == [real_length_id]
    assert result[0].content == b"drive bytes"
    assert result[0].filename == "Case Study - Zara Khan.docx"


def test_markaz_case_study_api_treats_401_as_a_documented_wall_not_an_error():
    """The endpoint returning 401 to automation must never surface as an
    exception or be mistaken for 'the submission doesn't exist'."""

    class FakeResponse:
        status_code = 401
        headers = {}
        content = b""

    def fake_get(url, timeout=10):
        return FakeResponse()

    ctx = _ctx()
    result = markaz_case_study_api(ctx, http_get=fake_get)  # must not raise
    assert result == []


def test_markaz_case_study_api_tries_both_word_and_excel():
    seen_urls = []

    class FakeResponse:
        status_code = 200
        headers = {"content-type": "application/octet-stream"}
        content = b"file bytes"

    def fake_get(url, timeout=10):
        seen_urls.append(url)
        return FakeResponse()

    ctx = _ctx()
    result = markaz_case_study_api(ctx, http_get=fake_get)
    assert len(result) == 2
    assert any(url.endswith("/4242/word") for url in seen_urls)
    assert any(url.endswith("/4242/excel") for url in seen_urls)


def test_mailbox_attachments_is_injectable_and_filters_by_application_id_in_filename():
    """Filtering must be by APPLICATION ID inside the filename (Rule 18), not
    by subject text or MIME type -- both files ride the same notification."""

    class FakeMessagePart:
        def __init__(self, filename, payload):
            self._filename = filename
            self._payload = payload

        def get_filename(self):
            return self._filename

        def get_payload(self, decode=False):
            return self._payload

    class FakeEmailMessage:
        def __init__(self, parts):
            self._parts = parts

        def walk(self):
            return iter(self._parts)

    import email as email_lib

    from webapp.services import submissions as submissions_module

    wanted = FakeEmailMessage([
        FakeMessagePart("case-study-4242-word-171.docx", b"wanted bytes"),
        FakeMessagePart("case-study-4242-excel-171.xlsx", b"wanted bytes 2"),
    ])
    other_app = FakeEmailMessage([
        FakeMessagePart("case-study-9999-word-171.docx", b"other candidate bytes"),
    ])

    class FakeIMAP:
        def __init__(self):
            self.logged_out = False

        def select(self, mailbox, readonly=True):
            return "OK", [b"1"]

        def search(self, charset, criteria):
            return "OK", [b"1 2"]

        def fetch(self, msg_id, parts):
            return "OK", [(b"1", b"placeholder")]

        def logout(self):
            self.logged_out = True

    fake_client = FakeIMAP()
    orig_from_bytes = email_lib.message_from_bytes
    call_order = {"n": 0}

    def fake_from_bytes(raw):
        call_order["n"] += 1
        return wanted if call_order["n"] == 1 else other_app

    email_lib.message_from_bytes = fake_from_bytes
    try:
        result = submissions_module.mailbox_attachments(_ctx(), imap_factory=lambda: fake_client)
    finally:
        email_lib.message_from_bytes = orig_from_bytes

    assert fake_client.logged_out is True
    origins = [r.origin for r in result]
    assert any("case-study-4242-word-171.docx" in o for o in origins)
    assert any("case-study-4242-excel-171.xlsx" in o for o in origins)
    assert not any("9999" in o for o in origins)


def test_mailbox_attachments_never_constructs_a_real_imap_client(monkeypatch):
    """Structural guarantee: without an injected imap_factory, this fetcher
    must go through _default_imap_client (lazy-imported), never silently
    reach the network some other way inside a test."""
    import imaplib

    def boom(*a, **kw):
        raise AssertionError("a unit test must never construct a real IMAP4_SSL")

    monkeypatch.setattr(imaplib, "IMAP4_SSL", boom)

    class FakeIMAP:
        def select(self, *a, **kw):
            return "OK", [b""]

        def search(self, *a, **kw):
            return "OK", [b""]

        def logout(self):
            pass

    from webapp.services import submissions as submissions_module

    result = submissions_module.mailbox_attachments(_ctx(), imap_factory=lambda: FakeIMAP())
    assert result == []


# ── Task 3-4 review fixes ────────────────────────────────────────────────
#
# FIX A: the 401 path (markaz_case_study_api, Rule 18) is now actually logged
# AND surfaced in corpus_for's own diagnostics, so a total-failure refusal
# message can distinguish "hit the documented 401 wall" from "found nothing".
# FIX B: the char floor no longer counts the injected "===== {origin} ====="
# provenance headers -- only the ACTUAL extracted text.


def test_corpus_for_names_the_401_wall_in_its_refusal_message():
    """The module docstring used to claim the 401 was 'logged and skipped' --
    it was skipped but never logged, and never surfaced anywhere a caller
    could see it. When every channel comes up empty, the refusal message
    must now say the 401 wall was hit, not just 'no source returned any
    content' -- exactly what the person fetching it by hand needs to know."""
    import functools

    from webapp.services.submissions import markaz_case_study_api

    class FakeResponse:
        status_code = 401
        headers = {}
        content = b""

    def fake_get(url, timeout=10):
        return FakeResponse()

    empty_fetcher = lambda ctx: []  # noqa: E731
    # Passed DIRECTLY (not wrapped in another lambda) so corpus_for's
    # _bind_diagnostics can see markaz_case_study_api's own `diagnostics`
    # parameter and inject its own `problems` list into it.
    api_fetcher = functools.partial(markaz_case_study_api, http_get=fake_get)

    with pytest.raises(SubmissionUnreadable) as exc_info:
        corpus_for(4242, context=_ctx(), fetchers=[empty_fetcher, api_fetcher])

    message = str(exc_info.value)
    assert "401" in message
    assert "case-study-file" in message
    assert "not a blocker" in message


def test_markaz_case_study_api_logs_the_401_at_info(caplog):
    """The docstring's claim, made real: a 401 is actually logged, not just
    silently skipped. Independent of corpus_for's diagnostics wiring."""
    import logging

    from webapp.services.submissions import markaz_case_study_api

    class FakeResponse:
        status_code = 401
        headers = {}
        content = b""

    def fake_get(url, timeout=10):
        return FakeResponse()

    with caplog.at_level(logging.INFO, logger="webapp.submissions"):
        result = markaz_case_study_api(_ctx(), http_get=fake_get)

    assert result == []  # still never raises, still never fabricates an attachment
    assert any("401" in r.message for r in caplog.records)


def test_markaz_case_study_api_diagnostics_stays_empty_when_no_list_is_passed():
    """A caller of markaz_case_study_api directly (not through corpus_for,
    e.g. every OTHER test in this file) never passes `diagnostics` -- the 401
    must still just log and return [], not raise for lack of somewhere to
    record itself."""
    from webapp.services.submissions import markaz_case_study_api

    class FakeResponse:
        status_code = 401
        headers = {}
        content = b""

    def fake_get(url, timeout=10):
        return FakeResponse()

    result = markaz_case_study_api(_ctx(), http_get=fake_get)  # no diagnostics kwarg
    assert result == []


def test_corpus_for_floor_excludes_provenance_headers_from_multiple_tiny_sources():
    """Several degraded sources each yielding only a couple of characters of
    REAL text must not pad past MIN_USABLE_CHARS on the injected
    '===== {origin} =====' provenance headers alone. Before this fix,
    `chars = len(combined)` counted the headers too; 20 sources with long
    descriptive origins easily clear 400 characters of header boilerplate
    while contributing only 40 characters of actual candidate text."""

    def make_fetcher(n):
        att = RawAttachment(
            origin=f"Gmail attachment: a rather long descriptive filename number {n}.txt",
            filename=f"note{n}.txt",
            content=b"hi",  # 2 REAL characters
        )
        return lambda ctx, _att=att: [_att]

    fetchers = [make_fetcher(n) for n in range(20)]

    # Sanity check on the OLD (buggy) measure: the headers alone exceed the
    # floor, which is exactly the defect this test guards against.
    header_only_estimate = sum(
        len(f"\n===== Gmail attachment: a rather long descriptive filename "
            f"number {n}.txt =====\nhi")
        for n in range(20)
    )
    assert header_only_estimate >= MIN_USABLE_CHARS

    with pytest.raises(SubmissionUnreadable) as exc_info:
        corpus_for(4242, context=_ctx(), fetchers=fetchers)
    # The refusal message reports the true (tiny) extracted-text count, not
    # the header-padded combined length.
    assert "got 40" in str(exc_info.value)


def test_corpus_for_chars_field_still_reflects_the_full_combined_text():
    """FIX B changes what GATES the refusal, not what `chars` reports on
    success -- the returned dict's `chars` stays len(text) (headers
    included), exactly as test_corpus_for_succeeds_and_records_every_source_
    actually_read already asserts. A single long source clears the new
    extracted-only floor easily, so this only needs to confirm the field
    didn't quietly change shape."""
    long_text = "Real candidate text about the case study. " * 20
    assert len(long_text) >= MIN_USABLE_CHARS
    att = RawAttachment(origin="Gmail attachment: real.docx", filename="real.txt",
                         content=long_text.encode("utf-8"))

    result = corpus_for(4242, context=_ctx(), fetchers=[lambda ctx: [att]])
    assert result["chars"] == len(result["text"])


# ── Task 6: the router (webapp/routers/case_studies.py) ────────────────────
#
# create_benchmark/approve_benchmark/score/get_evaluation are exercised
# directly against the route functions with a small in-memory fake Session
# (no live database, no real model call) -- unit tests of the router's own
# logic, mirroring test_values_scoring.py's _FakeSession. The dependency-
# identity gate test and the HTTP-level tests (test_case_studies_http.py)
# are what actually prove the auth wiring, not these.


class _FakeEvalResult:
    """Stands in for whatever a real SQLAlchemy Result needs to support at
    score()'s one raw-SQL call site: `.mappings().first()` for the Rule 0
    approved-benchmark-for-this-job lookup."""

    def __init__(self, mapping_row=None):
        self._mapping_row = mapping_row

    def mappings(self):
        return self

    def first(self):
        return self._mapping_row


_UNSET = object()


class _FakeCaseStudySession:
    """A minimal stand-in for a SQLAlchemy Session. No network, no engine.

    `add`/`get` keep EvalBenchmark/CaseStudyEvaluation rows in an in-memory
    dict keyed by `.id`, exactly like test_values_scoring.py's _FakeSession.

    `execute()` answers the ONE raw-SQL statement score() itself runs (the
    Rule 0 lookup, `FROM coco.eval_benchmarks`) one of two ways:
      - if the test set `approved_benchmark_row` explicitly, that row wins
        (the simple case: most tests just want "there is" / "there isn't"
        an approved benchmark);
      - otherwise it DERIVES the answer from whatever EvalBenchmark rows the
        test `add()`-ed, emulating the real query's
        `WHERE job_id = :job_id AND status = 'approved'
         ORDER BY qa_approved_at DESC NULLS LAST LIMIT 1` -- so a test can
        prove a DRAFT (or retired) benchmark is genuinely invisible to this
        lookup, not just assert the SQL string contains the word 'approved'.
    """

    def __init__(self):
        self.committed = 0
        self.rolled_back = 0
        self.execute_calls = []
        self.approved_benchmark_row = _UNSET
        self._store = {}

    def add(self, obj):
        # A real Session generates the primary key from the column's
        # Python-side `default=` at FLUSH time (inside commit()/add()
        # against a real engine) -- this fake never touches an engine, so it
        # emulates that one step explicitly for the two model types this
        # router constructs without an explicit `id=`, rather than silently
        # leaving `.id` as None (which a real flush never would).
        if getattr(obj, "id", None) is None:
            from webapp.models import CaseStudyEvaluation, EvalBenchmark, _benchmark_id, _evaluation_id

            if isinstance(obj, EvalBenchmark):
                obj.id = _benchmark_id()
            elif isinstance(obj, CaseStudyEvaluation):
                obj.id = _evaluation_id()
        self._store[getattr(obj, "id", None)] = obj

    def commit(self):
        self.committed += 1

    def rollback(self):
        self.rolled_back += 1

    def refresh(self, obj):
        pass

    def get(self, model, id_):
        return self._store.get(id_)

    def execute(self, stmt, params=None):
        sql = str(stmt)
        self.execute_calls.append((sql, params))
        if "FROM coco.eval_benchmarks" in sql:
            if self.approved_benchmark_row is not _UNSET:
                return _FakeEvalResult(mapping_row=self.approved_benchmark_row)
            from webapp.models import EvalBenchmark

            job_id = (params or {}).get("job_id")
            candidates = [
                b for b in self._store.values()
                if isinstance(b, EvalBenchmark) and b.job_id == job_id and b.status == "approved"
            ]
            candidates.sort(key=lambda b: b.qa_approved_at or dt.datetime.min, reverse=True)
            row = {"id": candidates[0].id, "body": candidates[0].body} if candidates else None
            return _FakeEvalResult(mapping_row=row)
        return _FakeEvalResult()


def _fake_case_study_app_row(application_id=555):
    return {
        "application_id": application_id,
        "first_name": "Zara",
        "last_name": "Khan",
        "job_title": "Growth Manager",
        "job_pk": 7,
    }


# --- benchmarks: create -----------------------------------------------------


def test_create_benchmark_starts_as_draft_with_no_qa_fields_set():
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyBenchmarkCreateRequest

    db = _FakeCaseStudySession()
    body = CaseStudyBenchmarkCreateRequest(job_id=7, title="Growth flywheel", body="the case text")
    out = router_mod.create_benchmark(body, db, {"id": "appuser-editor"})

    assert out["job_id"] == 7
    assert out["title"] == "Growth flywheel"
    assert out["status"] == "draft"
    # The create body carries no qa_approved_by/at fields at all (see
    # CaseStudyBenchmarkCreateRequest) -- there is no way for a client to
    # hand this endpoint an already-approved benchmark.
    assert out["qa_approved_by"] is None
    assert out["qa_approved_at"] is None
    assert db.committed == 1


# --- benchmarks: approve -----------------------------------------------------


def test_approve_benchmark_sets_qa_fields_and_flips_status_to_approved():
    import webapp.routers.case_studies as router_mod
    from webapp.models import EvalBenchmark

    db = _FakeCaseStudySession()
    db._store["benchmark-1"] = EvalBenchmark(
        id="benchmark-1", job_id=7, kind="case_study", title="t", body="b",
        created_by="appuser-editor", status="draft",
    )

    out = router_mod.approve_benchmark(
        "benchmark-1", db, {"id": "appuser-approver", "email": "ayesha.khan@taleemabad.com"}
    )

    assert out["status"] == "approved"
    assert out["qa_approved_by"] == "appuser-approver ayesha.khan@taleemabad.com"
    assert out["qa_approved_at"] is not None
    assert db.committed == 1


def test_approve_benchmark_refuses_a_second_approval():
    import webapp.routers.case_studies as router_mod
    from webapp.models import EvalBenchmark

    db = _FakeCaseStudySession()
    db._store["benchmark-2"] = EvalBenchmark(
        id="benchmark-2", job_id=7, kind="case_study", title="t", body="b",
        created_by="x", status="approved", qa_approved_by="someone else",
    )

    with pytest.raises(HTTPException) as exc_info:
        router_mod.approve_benchmark("benchmark-2", db, {"id": "appuser-approver"})
    assert exc_info.value.status_code == 409
    assert db.committed == 0


def test_approve_benchmark_refuses_a_retired_benchmark():
    import webapp.routers.case_studies as router_mod
    from webapp.models import EvalBenchmark

    db = _FakeCaseStudySession()
    db._store["benchmark-3"] = EvalBenchmark(
        id="benchmark-3", job_id=7, kind="case_study", title="t", body="b",
        created_by="x", status="retired",
    )

    with pytest.raises(HTTPException) as exc_info:
        router_mod.approve_benchmark("benchmark-3", db, {"id": "appuser-approver"})
    assert exc_info.value.status_code == 409
    assert db.committed == 0


def test_approve_benchmark_404_on_a_missing_id():
    import webapp.routers.case_studies as router_mod

    db = _FakeCaseStudySession()
    with pytest.raises(HTTPException) as exc_info:
        router_mod.approve_benchmark("benchmark-nope", db, {"id": "appuser-approver"})
    assert exc_info.value.status_code == 404


# --- score: RULE 0 -----------------------------------------------------------


def test_score_refuses_with_409_when_no_approved_benchmark_exists_for_the_job(monkeypatch):
    """RULE 0. THE POINT OF THIS TASK. No approved benchmark row for the
    application's job -> 409, BEFORE the submission is ever fetched or
    scored -- corpus_for and score_submission must never even be called."""
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyScoreRequest

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )
    corpus_calls = []
    monkeypatch.setattr(
        router_mod, "corpus_for",
        lambda app_id, *, db=None: corpus_calls.append(app_id) or {"text": "x", "sources": [], "chars": 1, "usable": True},
    )
    score_calls = []
    monkeypatch.setattr(router_mod, "score_submission", lambda **kw: score_calls.append(kw))

    db = _FakeCaseStudySession()
    db.approved_benchmark_row = None  # no approved benchmark for job 7

    with pytest.raises(HTTPException) as exc_info:
        router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})

    assert exc_info.value.status_code == 409
    assert "Rule 0" in exc_info.value.detail
    assert "7" in exc_info.value.detail  # names the job that has no benchmark
    assert corpus_calls == []
    assert score_calls == []
    assert db.committed == 0


def test_score_refuses_409_for_a_draft_benchmark_too(monkeypatch):
    """A DRAFT benchmark for the right job must be exactly as invisible to
    Rule 0 as no benchmark at all -- 409, not a silent pass. Uses the fake
    session's STORE-DERIVED lookup (not a pre-set `approved_benchmark_row`),
    so this genuinely exercises "only status='approved' rows are found",
    not merely a string in the SQL constant."""
    import webapp.routers.case_studies as router_mod
    from webapp.models import EvalBenchmark
    from webapp.schemas import CaseStudyScoreRequest

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )
    score_calls = []
    monkeypatch.setattr(router_mod, "score_submission", lambda **kw: score_calls.append(kw))

    db = _FakeCaseStudySession()
    db._store["benchmark-draft"] = EvalBenchmark(
        id="benchmark-draft", job_id=7, kind="case_study", title="t", body="b",
        created_by="x", status="draft",
    )
    # approved_benchmark_row is left at _UNSET, so execute() derives the
    # answer from the store above -- a real WHERE status='approved' semantic.

    with pytest.raises(HTTPException) as exc_info:
        router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})

    assert exc_info.value.status_code == 409
    assert score_calls == []


def test_score_finds_an_approved_benchmark_via_the_store_derived_lookup_and_ignores_a_draft_sibling(
    monkeypatch,
):
    """The positive case for the same store-derived lookup: an approved
    benchmark for the job is found and used even when a draft sibling for
    the SAME job also exists in the store -- proving the filter is real,
    not just "any row present"."""
    import webapp.routers.case_studies as router_mod
    from webapp.models import EvalBenchmark
    from webapp.schemas import CaseStudyScoreRequest

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )
    monkeypatch.setattr(
        router_mod, "corpus_for",
        lambda app_id, *, db=None: {
            "text": "submission text " * 50, "sources": ["Gmail attachment: case.docx"],
            "chars": 800, "usable": True,
        },
    )
    captured = {}

    def fake_score_submission(*, corpus, benchmark_body, candidate_name, role):
        captured["benchmark_body"] = benchmark_body
        return {
            "scores": _all(4), "evidence": _good_evidence(), "flags": [],
            "total": 80.0, "band": "strong_yes", "model": "test-model",
        }

    monkeypatch.setattr(router_mod, "score_submission", fake_score_submission)

    db = _FakeCaseStudySession()
    db._store["benchmark-draft-sibling"] = EvalBenchmark(
        id="benchmark-draft-sibling", job_id=7, kind="case_study", title="t",
        body="the wrong, unapproved body", created_by="x", status="draft",
    )
    db._store["benchmark-approved"] = EvalBenchmark(
        id="benchmark-approved", job_id=7, kind="case_study", title="t",
        body="the right, approved body", created_by="x", status="approved",
        qa_approved_by="appuser-approver", qa_approved_at=dt.datetime(2026, 9, 20, tzinfo=dt.timezone.utc),
    )

    out = router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})

    assert out["benchmark_id"] == "benchmark-approved"
    assert captured["benchmark_body"] == "the right, approved body"


# --- score: the rest of the happy/refusal paths -----------------------------


def test_score_404_when_the_application_is_not_found(monkeypatch):
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyScoreRequest

    monkeypatch.setattr(router_mod.reads, "get_application", lambda db, app_id: None)
    db = _FakeCaseStudySession()

    with pytest.raises(HTTPException) as exc_info:
        router_mod.score(CaseStudyScoreRequest(application_id=9999), db, {"id": "appuser-editor"})
    assert exc_info.value.status_code == 404


def test_score_refuses_422_when_the_submission_is_unreadable(monkeypatch):
    """An unreadable submission is refused, never scored as a weak one
    (Rule 26 / cv_text's refusal pattern) -- score_submission must never
    even be called."""
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyScoreRequest
    from webapp.services.submissions import SubmissionUnreadable as _SU

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )

    def fake_corpus_for(app_id, *, db=None):
        raise _SU(f"no usable case-study text for application {app_id}")

    monkeypatch.setattr(router_mod, "corpus_for", fake_corpus_for)
    score_calls = []
    monkeypatch.setattr(router_mod, "score_submission", lambda **kw: score_calls.append(kw))

    db = _FakeCaseStudySession()
    db.approved_benchmark_row = {"id": "benchmark-1", "body": "b"}

    with pytest.raises(HTTPException) as exc_info:
        router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})

    assert exc_info.value.status_code == 422
    assert score_calls == []
    assert db.committed == 0


def test_score_succeeds_and_persists_the_evaluation_with_the_benchmark_actually_used(monkeypatch):
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyScoreRequest

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )
    monkeypatch.setattr(
        router_mod, "corpus_for",
        lambda app_id, *, db=None: {
            "text": "submission text " * 50, "sources": ["Gmail attachment: case.docx"],
            "chars": 800, "usable": True,
        },
    )
    captured = {}

    def fake_score_submission(*, corpus, benchmark_body, candidate_name, role):
        captured.update(corpus=corpus, benchmark_body=benchmark_body,
                         candidate_name=candidate_name, role=role)
        return {
            "scores": _all(4), "evidence": _good_evidence(), "flags": [],
            "total": 80.0, "band": "strong_yes", "model": "test-model",
        }

    monkeypatch.setattr(router_mod, "score_submission", fake_score_submission)

    db = _FakeCaseStudySession()
    db.approved_benchmark_row = {"id": "benchmark-approved-1", "body": "the answer key body"}

    out = router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})

    assert out["benchmark_id"] == "benchmark-approved-1"
    assert out["total"] == 80.0
    assert out["band"] == "strong_yes"
    assert out["candidate_name"] == "Zara Khan"
    assert out["job_id"] == 7
    assert out["sources"] == ["Gmail attachment: case.docx"]
    assert captured["benchmark_body"] == "the answer key body"
    assert captured["candidate_name"] == "Zara Khan"
    assert captured["role"] == "Growth Manager"
    assert db.committed == 1


def test_score_maps_a_disqualifying_flag_all_the_way_through_to_the_persisted_row(monkeypatch):
    """The total/band are never recomputed here -- taken AS-IS from
    score_submission's own return value (which itself already refused to
    trust the model, per Task 5). A disqualified band with a high total
    must survive the round trip unchanged."""
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyScoreRequest

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )
    monkeypatch.setattr(
        router_mod, "corpus_for",
        lambda app_id, *, db=None: {
            "text": "submission text " * 50, "sources": ["Gmail attachment: case.docx"],
            "chars": 800, "usable": True,
        },
    )
    monkeypatch.setattr(
        router_mod, "score_submission",
        lambda **kw: {
            "scores": _all(5), "evidence": _good_evidence(), "flags": ["fabricated_data"],
            "total": 100.0, "band": "disqualified", "model": "test-model",
        },
    )

    db = _FakeCaseStudySession()
    db.approved_benchmark_row = {"id": "benchmark-1", "body": "b"}

    out = router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})
    assert out["total"] == 100.0
    assert out["band"] == "disqualified"
    assert out["flags"] == ["fabricated_data"]


def test_score_maps_drafting_unavailable_to_503_not_500(monkeypatch):
    import webapp.routers.case_studies as router_mod
    from webapp.schemas import CaseStudyScoreRequest
    from webapp.services.drafting import DraftingUnavailable

    monkeypatch.setattr(
        router_mod.reads, "get_application", lambda db, app_id: _fake_case_study_app_row(app_id)
    )
    monkeypatch.setattr(
        router_mod, "corpus_for",
        lambda app_id, *, db=None: {
            "text": "submission text " * 50, "sources": ["Gmail attachment: case.docx"],
            "chars": 800, "usable": True,
        },
    )

    def boom(**kw):
        raise DraftingUnavailable("no credential configured")

    monkeypatch.setattr(router_mod, "score_submission", boom)

    db = _FakeCaseStudySession()
    db.approved_benchmark_row = {"id": "benchmark-1", "body": "b"}

    with pytest.raises(HTTPException) as exc_info:
        router_mod.score(CaseStudyScoreRequest(application_id=555), db, {"id": "appuser-editor"})
    assert exc_info.value.status_code == 503


# --- get_evaluation -----------------------------------------------------------


def test_get_evaluation_404_when_missing():
    import webapp.routers.case_studies as router_mod

    db = _FakeCaseStudySession()
    with pytest.raises(HTTPException) as exc_info:
        router_mod.get_evaluation("cse-missing", db, {"id": "appuser-editor"})
    assert exc_info.value.status_code == 404


def test_get_evaluation_returns_the_persisted_row_exactly():
    import webapp.routers.case_studies as router_mod
    from webapp.models import CaseStudyEvaluation

    db = _FakeCaseStudySession()
    db._store["cse-1"] = CaseStudyEvaluation(
        id="cse-1", application_id=555, job_id=7, benchmark_id="benchmark-1",
        candidate_name="Zara Khan", role="Growth Manager",
        scores=_all(3), evidence=_good_evidence(), flags=[],
        total=60.0, band="yes", model_name="test-model",
        sources=["Gmail attachment: case.docx"], created_by="appuser-editor",
    )

    out = router_mod.get_evaluation("cse-1", db, {"id": "appuser-editor"})
    assert out["id"] == "cse-1"
    assert out["total"] == 60.0
    assert out["band"] == "yes"
    assert out["model"] == "test-model"  # read off model_name, the renamed attribute
    assert out["benchmark_id"] == "benchmark-1"


# --- Rule 0 gate identity -----------------------------------------------------


def _flatten_case_study_routes(routes):
    """See webapp/tests/test_values_scoring.py::_flatten_routes and
    webapp/tests/test_nugget_reads.py::_flatten_routes -- this FastAPI build
    (0.137.1) wraps every `include_router()` call in an `_IncludedRouter`
    shim with no `.path`/`.methods` of its own; the real `APIRoute` objects
    live on `original_router.routes`. Recurse through those shims so route
    introspection sees a flat list."""
    flat = []
    for r in routes:
        if hasattr(r, "path"):
            flat.append(r)
        nested = getattr(r, "original_router", None)
        if nested is not None:
            flat.extend(_flatten_case_study_routes(nested.routes))
    return flat


def _case_study_route(path, method):
    from webapp.main import app

    for r in _flatten_case_study_routes(app.routes):
        if r.path == path and method in getattr(r, "methods", set()):
            return r
    raise AssertionError(f"route not found: {method} {path}")


def test_the_case_study_routes_are_gated_on_the_declared_dependency_not_the_docstring():
    """Was (in Phase 2, before review caught it):
    `"require_approver" in src` against the whole module's SOURCE TEXT --
    this router's own docstring names require_approver/require_editor
    several times, so that assertion would pass even if a route were wired
    to the wrong gate. Do NOT repeat that mistake here.

    This inspects the route object FastAPI actually built --
    `route.dependant.dependencies` -- and checks the resolved callable BY
    IDENTITY. `require_editor`/`require_approver` are each built once at
    import time by `_require_role(...)` in webapp/deps.py, so `is`/`in` over
    a set of the actual callables is a valid, exact check, and
    `get_current_user` (their own sub-dependency) never appears in a route's
    own top-level dependency list -- only the gate wrapper does.
    """
    import webapp.deps as deps

    create_benchmark = _case_study_route("/api/case-studies/benchmarks", "POST")
    approve_benchmark = _case_study_route(
        "/api/case-studies/benchmarks/{benchmark_id}/approve", "POST"
    )
    score_route = _case_study_route("/api/case-studies/score", "POST")
    get_evaluation = _case_study_route("/api/case-studies/evaluations/{evaluation_id}", "GET")

    for route, expected in (
        (create_benchmark, deps.require_editor),
        (approve_benchmark, deps.require_approver),
        (score_route, deps.require_editor),
        (get_evaluation, deps.require_editor),
    ):
        calls = {dep.call for dep in route.dependant.dependencies}
        assert expected in calls, f"{route.path} [{route.methods}] missing {expected}"
        other = deps.require_approver if expected is deps.require_editor else deps.require_editor
        assert other not in calls, f"{route.path} [{route.methods}] wrongly gated on {other}"
        # get_current_user is a SUB-dependency of require_editor/require_approver,
        # never wired directly on any of these four routes (a bare signed-in
        # user is not enough for any of them).
        assert deps.get_current_user not in calls
