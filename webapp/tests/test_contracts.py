"""Contract drafting rules, and the ones that would do real harm if wrong.

Run: python -m pytest webapp/tests/test_contracts.py -v

The two that matter most:
  - A VOLUNTEER FELLOW NEVER RECEIVES A CONTRACT. Sending an unpaid person an
    employment contract creates an obligation nobody agreed to.
  - A MISSING MASTER IS A BLOCKER, never a substitution from another entity.
    These are approved legal documents; the wrong letterhead is a real defect.

Field discovery is checked against the REAL masters where they are present on
disk, because the whole design rests on the claim that a fill field is a
yellow-highlighted run. Skipped when the folder is absent (it is gitignored),
so this suite still runs on a fresh checkout.
"""

from __future__ import annotations

import pathlib

import pytest

from webapp.services import contracts as c

REPO = pathlib.Path(__file__).resolve().parents[2]
MASTERS = REPO / "Contracts"

pytest_plugins: list[str] = []


def _master(rel: str) -> bytes:
    path = MASTERS / rel
    if not path.is_file():
        pytest.skip(f"master not on this machine: {rel}")
    return path.read_bytes()


# --------------------------------------------------------------------------
# Package rules
# --------------------------------------------------------------------------


@pytest.mark.parametrize("entity", c.ENTITIES)
def test_a_volunteer_fellow_gets_an_nda_and_never_a_contract(entity):
    docs = c.documents_for(c.VOLUNTEER_FELLOW, entity)
    assert docs == [c.FELLOW_NDA]
    assert c.FELLOW_CONTRACT not in docs
    assert not any("contract" in d for d in docs)


@pytest.mark.parametrize("entity", c.ENTITIES)
def test_a_paid_fellow_gets_both_documents(entity):
    assert c.documents_for(c.PAID_FELLOW, entity) == [c.FELLOW_CONTRACT, c.FELLOW_NDA]


@pytest.mark.parametrize("entity", c.ENTITIES)
def test_an_unpaid_to_paid_fellow_gets_the_contract_only(entity):
    """The NDA was signed when they started. A second one implies the first
    did not count."""
    assert c.documents_for(c.FELLOW_TO_PAID, entity) == [c.FELLOW_CONTRACT]


def test_a_same_team_promotion_is_an_addendum_and_a_team_move_is_a_contract():
    assert c.documents_for(c.PROMOTION_SAME_TEAM, c.OPL) == [c.ADDENDUM]
    assert c.documents_for(c.TEAM_MOVE, c.OPL) == [c.PERMANENT_CONTRACT]


def test_every_engagement_has_a_document_rule():
    """A new engagement type with no rule must fail loudly, not return []."""
    for engagement in c.ENGAGEMENTS:
        docs = c.documents_for(engagement, c.OPL)
        assert docs, engagement


def test_an_unknown_engagement_or_entity_raises():
    with pytest.raises(c.ContractError):
        c.documents_for("freelance", c.OPL)
    with pytest.raises(c.ContractError):
        c.documents_for(c.PAID_FELLOW, "Orenda")


# --------------------------------------------------------------------------
# Masters: routed, or refused
# --------------------------------------------------------------------------


def test_the_nda_and_addendum_masters_serve_every_entity():
    """One Orenda NDA covers OPL, OWT, NIETE and Inc. (Ayesha 2026-08-12)."""
    for doc_type in (c.PERMANENT_NDA, c.FELLOW_NDA, c.ADDENDUM):
        paths = {c.master_for(e, doc_type) for e in c.ENTITIES}
        assert len(paths) == 1 and None not in paths, doc_type


def test_a_missing_master_blocks_the_plan_rather_than_substituting():
    """NIETE has no permanent full-time master. The plan must say so."""
    assert c.master_for(c.NIETE, c.PERMANENT_CONTRACT) is None
    p = c.plan(c.PERMANENT_HIRE, c.NIETE)
    assert p["blockers"], "a missing master must be reported"
    assert "no approved master" in p["blockers"][0]


def test_a_complete_plan_has_no_blockers():
    p = c.plan(c.PAID_FELLOW, c.OPL)
    assert p["blockers"] == []
    assert [d["doc_type"] for d in p["documents"]] == [c.FELLOW_CONTRACT, c.FELLOW_NDA]


def test_every_plan_carries_the_visual_proof_caveat():
    """Rule 14: structural checks are not visual proof, said every time
    rather than remembered."""
    for engagement in c.ENGAGEMENTS:
        p = c.plan(engagement, c.OPL)
        assert "never appearance" in p["caveat"]


def test_owt_project_based_is_refused_because_there_is_only_owt_full_time():
    assert c.master_for(c.OWT, c.PROJECT_CONTRACT) is None


# --------------------------------------------------------------------------
# Field discovery, against the real masters
# --------------------------------------------------------------------------

# Highlighted fill fields PLUS the placeholders nobody highlighted. The NIETE
# contract has 15 highlighted and 5 plain-text (the acceptance joining date and
# the four salary lines), so 20. The NDAs are fully highlighted, so 4 each.
REAL = {
    "Promotion/Template - NDA Full Time Permanent Employee.docx": 4,
    "Fellow/Template - NDA Fellow Employee.docx": 4,
    "NIETE/NIETE - Project-based Employment Contract.docx": 20,
}


@pytest.mark.parametrize("rel,expected", sorted(REAL.items()))
def test_fields_are_discovered_from_the_real_masters(rel, expected):
    """The design rests on a fill field being a highlighted run. If a master
    is re-issued with a different number of fields, this fails, which is the
    point: the field list must follow the document."""
    fields = c.discover_fields(_master(rel))
    assert len(fields) == expected, [f["placeholder"] for f in fields]


def test_the_permanent_nda_collapses_to_three_inputs():
    """EMPLOYEE NAME appears twice and means the same person, so it is typed
    once and written to both places."""
    groups = c.group_fields(c.discover_fields(
        _master("Promotion/Template - NDA Full Time Permanent Employee.docx")))
    keys = [g["key"] for g in groups]
    assert keys == ["EMPLOYEE NAME", "JOINING DATE", "CURRENT DATE"]
    name = groups[0]
    assert len(name["indexes"]) == 2
    # It writes into two different sentences, so the UI must say so.
    assert name["spans_contexts"] is True


def test_opaque_placeholders_are_never_merged():
    """"XYZ" is a duration in one sentence and a CNIC in another. Merging them
    would put a national ID number where a contract term belongs."""
    groups = c.group_fields(c.discover_fields(
        _master("NIETE/NIETE - Project-based Employment Contract.docx")))
    opaque = [g for g in groups if g["opaque"]]
    assert len(opaque) >= 3
    for g in opaque:
        assert len(g["indexes"]) == 1, f"{g['key']} was merged with another field"
        assert g["contexts"][0], "an opaque field must carry its sentence"


@pytest.mark.parametrize("text", ["XYZ", "xyz", "X Y Z", " xyz ", "x-y-z"])
def test_bare_placeholders_are_recognised_as_opaque(text):
    assert c.is_opaque(text)


@pytest.mark.parametrize(
    "text", ["EMPLOYEE NAME", "JOINING DATE", "CURRENT DATE", "Direct Report to:"]
)
def test_self_describing_placeholders_are_not_opaque(text):
    assert not c.is_opaque(text)


def test_the_opaque_check_bites():
    """Without this, an is_opaque that returned False for everything would
    make the merge tests pass while merging nothing."""
    assert c.is_opaque("XYZ") and not c.is_opaque("EMPLOYEE NAME")


# --------------------------------------------------------------------------
# A contract never ships a blank
# --------------------------------------------------------------------------


def test_every_input_must_be_filled_before_building():
    """Skill 07 Rule 6: never invent a field, leave it visibly open and ask."""
    groups = c.group_fields(c.discover_fields(
        _master("Promotion/Template - NDA Full Time Permanent Employee.docx")))
    assert len(c.missing_values(groups, {})) == len(groups)
    filled = {g["key"]: "x" for g in groups}
    assert c.missing_values(groups, filled) == []
    # Whitespace is not a value.
    filled["JOINING DATE"] = "   "
    assert c.missing_values(groups, filled) == ["JOINING DATE"]


def test_one_typed_value_reaches_every_place_it_fills():
    groups = c.group_fields(c.discover_fields(
        _master("Promotion/Template - NDA Full Time Permanent Employee.docx")))
    by_index = c.values_by_index(groups, {
        "EMPLOYEE NAME": "Hajra Noor", "JOINING DATE": "1 October 2026",
        "CURRENT DATE": "25 September 2026",
    })
    assert by_index[0] == "Hajra Noor"
    assert by_index[3] == "Hajra Noor"   # the printed-name line
    assert by_index[1] == "1 October 2026"
    assert by_index[2] == "25 September 2026"


# --------------------------------------------------------------------------
# The placeholders nobody highlighted
# --------------------------------------------------------------------------

# On the first real build of a NIETE contract, filling every HIGHLIGHTED field
# still left five placeholders printed in the document: the acceptance-line
# joining date, and all four salary lines, because only ONE of
# "Total Earnings PKR XYZ / Base Salary: PKR XYZ / Medical: PKR XYZ /
# Others: PKR XYZ" carries highlighting. The validator called them WARNINGs and
# the wrapper reported passed=True. A contract reading "PKR XYZ" would have
# been offered for download.

NIETE_MASTER = "NIETE/NIETE - Project-based Employment Contract.docx"


def test_the_unhighlighted_salary_lines_are_discovered_as_fields():
    from webapp.services import contracts as svc

    fields = svc.discover_fields(_master(NIETE_MASTER))
    text_fields = [f for f in fields if f.get("text_token")]
    assert len(text_fields) == 5, [f["context"][:60] for f in text_fields]
    salary = [f for f in text_fields if "PKR" in f["context"]]
    assert len(salary) == 4, "all four salary lines must be fillable"


def test_a_highlighted_placeholder_is_not_also_listed_as_a_text_field():
    """Otherwise the same value is asked for twice and the count is
    meaningless."""
    from webapp.services import contracts as svc

    fields = svc.discover_fields(_master(NIETE_MASTER))
    assert len(fields) == 20, len(fields)          # 15 highlighted + 5 plain
    assert len([f for f in fields if not f.get("text_token")]) == 15


def test_a_fully_filled_contract_has_no_placeholders_left():
    from webapp.services import contract_build as build
    from webapp.services import contracts as svc

    data = _master(NIETE_MASTER)
    fields = svc.discover_fields(data)
    groups = svc.group_fields(fields)
    out = build.fill(
        data, svc.values_by_index(groups, {g["key"]: "FILLED" for g in groups}),
        fields=fields,
    )
    assert svc.unresolved_placeholders(out) == []


def test_an_unfilled_salary_line_is_a_hard_block_not_a_warning():
    """Proof the gate bites, against a deliberately broken document. A gate
    that has only ever passed proves nothing."""
    from webapp.services import contract_build as build
    from webapp.services import contracts as svc

    data = _master(NIETE_MASTER)
    fields = svc.discover_fields(data)
    groups = svc.group_fields(fields)
    keep = {g["key"]: "FILLED" for g in groups
            if not (g["opaque"] and "PKR" in g["contexts"][0])}
    by_index = {k: v for k, v in svc.values_by_index(groups, keep).items() if v}
    broken = build.fill(data, by_index, fields=fields)

    report = build.validate(broken, svc.PROJECT_CONTRACT)
    assert report["passed"] is False
    hard = [f for f in report["findings"] if f["severity"] == "HARD_BLOCK"]
    assert len(hard) == 4, hard
    assert all("PKR" in f["message"] for f in hard)


def test_the_nda_still_builds_and_passes():
    from webapp.services import contract_build as build
    from webapp.services import contracts as svc

    data = _master("Promotion/Template - NDA Full Time Permanent Employee.docx")
    fields = svc.discover_fields(data)
    groups = svc.group_fields(fields)
    out = build.fill(data, svc.values_by_index(groups, {
        "EMPLOYEE NAME": "Hajra Noor",
        "JOINING DATE": "1 October 2026",
        "CURRENT DATE": "25 September 2026",
    }), fields=fields)
    assert svc.unresolved_placeholders(out) == []
    text = "\n".join(p.text for p in _open(out).paragraphs)
    assert "Hajra Noor" in text and "EMPLOYEE NAME" not in text


def _open(data: bytes):
    import io

    from docx import Document

    return Document(io.BytesIO(data))


def test_values_for_fields_that_do_not_exist_are_refused():
    """A master re-issued with fewer fields must fail loudly rather than
    shifting every value one place along."""
    from webapp.services import contract_build as build
    from webapp.services import contracts as svc

    data = _master("Promotion/Template - NDA Full Time Permanent Employee.docx")
    with pytest.raises(build.BuildError, match="re-issued"):
        build.fill(data, {0: "a", 1: "b", 2: "c", 3: "d", 99: "nowhere"})


def test_the_validator_type_is_never_guessed():
    """contract_docx_eval --type defaults to `fellow`, and a project contract
    checked as a fellow one throws 11 false hard blocks (Rule 19)."""
    from webapp.services import contract_build as build
    from webapp.services import contracts as svc

    assert build._VALIDATOR_TYPE[svc.PROJECT_CONTRACT] == "project"
    assert build._VALIDATOR_TYPE[svc.FELLOW_CONTRACT] == "fellow"
    with pytest.raises(build.BuildError, match="No validator profile"):
        build.validate(b"", "something_new")


def test_the_generated_filename_carries_no_identity_number():
    from webapp.services import contract_build as build
    from webapp.services import contracts as svc

    name = build.filename_for(svc.PERMANENT_NDA, "Hajra Noor")
    assert name == "Permanent Employee NDA - Hajra Noor.docx"
    assert "/" not in name and "\\" not in name
