"""Extract plain text from a candidate's stored CV.

Markaz stores the uploaded resume in `candidates.resume_data` as a base64 TEXT
column (not bytea), with `resume_mime_type` and `resume_file_name` alongside.

This module is the ONLY place the webapp turns that blob into text. A CV-stage
rejection must be grounded in the candidate's actual CV (Skill 01, step 1:
"pull resume_data from candidates table"), so the drafting path needs real text
or it needs to refuse — never a silent empty string, which is what lets a model
invent the CV it was never shown.
"""

from __future__ import annotations

import base64
import binascii
import io
import re
from typing import Optional

# Below this, whatever we extracted is not a readable CV: an image-only scan, a
# parse that produced page furniture and nothing else, or a corrupt upload. The
# skill's own floor is 10,000 chars for a *full* read; this is the far lower bar
# that separates "text" from "no text", and the caller decides what to do.
MIN_USABLE_CHARS = 400


class CVUnreadable(Exception):
    """The CV exists but no usable text could be extracted from it."""


def _decode(resume_data: str | bytes) -> bytes:
    if isinstance(resume_data, bytes):
        # Already binary (defensive: column is TEXT today, could change).
        if resume_data[:5] == b"%PDF-" or resume_data[:2] == b"PK":
            return resume_data
        resume_data = resume_data.decode("ascii", "ignore")
    cleaned = re.sub(r"\s+", "", resume_data)
    try:
        return base64.b64decode(cleaned, validate=False)
    except (binascii.Error, ValueError) as exc:
        raise CVUnreadable(f"resume_data is not valid base64: {exc}") from exc


def _from_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(raw: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    # Many CVs lay everything out in tables; paragraphs alone miss all of it.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _despace(text: str) -> str:
    """Repair letter-spaced PDF extraction.

    Some PDFs (font/encoding dependent) come out of pypdf with a space between
    every character and a DOUBLE space between words:

        "W O R K  E X P E R I E N C E"  ->  "WORK EXPERIENCE"

    Measured on the Job-42 cohort this hit 12 of 98 CVs. Left alone it is far
    worse than a failed parse, because the text looks substantial (20,089 chars
    for one candidate) and sails past any length check, while carrying almost no
    readable words: the drafter is handed noise and told it is the CV.

    Must run BEFORE whitespace is collapsed — collapsing destroys the double
    space that separates the words.
    """
    out = []
    for line in text.replace("\r\n", "\n").split("\n"):
        tokens = [t for t in line.split(" ") if t]
        if len(tokens) >= 6 and sum(1 for t in tokens if len(t) == 1) / len(tokens) > 0.6:
            words = [re.sub(r"\s+", "", w) for w in re.split(r" {2,}", line.strip())]
            out.append(" ".join(w for w in words if w))
        else:
            out.append(line)
    return "\n".join(out)


def _tidy(text: str) -> str:
    # Before any whitespace is collapsed: the double space is the only thing
    # separating words in a letter-spaced extraction.
    text = _despace(text.replace("\xa0", " "))
    text = text.replace("\r\n", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def extract(
    resume_data: Optional[str],
    *,
    mime_type: Optional[str] = None,
    file_name: Optional[str] = None,
) -> str:
    """Return readable CV text, or raise CVUnreadable.

    Sniffs the real file signature rather than trusting `resume_mime_type` —
    Markaz records are inconsistent, and a .docx labelled application/pdf is
    common enough to matter.
    """
    if not resume_data:
        raise CVUnreadable("no resume on file for this candidate")

    raw = _decode(resume_data)
    hint = f"{mime_type or ''} {file_name or ''}".lower()

    if raw[:5] == b"%PDF-":
        order = (_from_pdf, _from_docx)
    elif raw[:2] == b"PK":  # zip container: docx/odt
        order = (_from_docx, _from_pdf)
    elif "pdf" in hint:
        order = (_from_pdf, _from_docx)
    else:
        order = (_from_docx, _from_pdf)

    errors = []
    for parser in order:
        try:
            text = _tidy(parser(raw))
        except Exception as exc:  # noqa: BLE001 - any parser failure is "try the next one"
            errors.append(f"{parser.__name__}: {type(exc).__name__}: {exc}")
            continue
        if len(text) >= MIN_USABLE_CHARS:
            return text
        errors.append(f"{parser.__name__}: only {len(text)} chars of text")

    name = file_name or "the uploaded file"
    raise CVUnreadable(
        f"could not extract usable text from {name} "
        f"(needs {MIN_USABLE_CHARS}+ chars). Tried: {'; '.join(errors)}. "
        "If this is a scanned or image-only CV it must be read by hand."
    )
